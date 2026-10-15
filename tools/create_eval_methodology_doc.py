#!/usr/bin/env python3
"""Create a concise Vietnamese Word note on the current MPPI/SAC evaluation."""

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "reports" / "so_do_danh_gia_mppi_sac_hai_map_warehouse.docx"
DIAGRAM = Path("/tmp/mppi_sac_evaluation_flow.png")
FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"


def line(draw, xy, fill="#5E7186", width=4):
    draw.line(xy, fill=fill, width=width, joint="curve")


def arrow(draw, start, end, fill="#5E7186"):
    line(draw, [start, end], fill=fill)
    import math

    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    spread = 0.45
    length = 16
    tip1 = (end[0] - length * math.cos(angle - spread), end[1] - length * math.sin(angle - spread))
    tip2 = (end[0] - length * math.cos(angle + spread), end[1] - length * math.sin(angle + spread))
    draw.polygon([end, tip1, tip2], fill=fill)


def box(draw, rect, title, lines, fill, font_title, font_body):
    draw.rounded_rectangle(rect, radius=17, fill=fill, outline="#C9D4DF", width=2)
    x1, y1, x2, _ = rect
    draw.text((x1 + 22, y1 + 20), title, font=font_title, fill="#122334")
    y = y1 + 72
    for item in lines:
        draw.text((x1 + 22, y), item, font=font_body, fill="#294258")
        y += 38


def make_diagram():
    im = Image.new("RGB", (1800, 830), "#FFFFFF")
    d = ImageDraw.Draw(im)
    ft = ImageFont.truetype(FONT_BOLD, 34)
    fb = ImageFont.truetype(FONT, 26)
    small = ImageFont.truetype(FONT, 23)
    box(d, (35, 65, 485, 265), "Kịch bản mô phỏng", ["Map warehouse có vật cản tĩnh", "Vật cản động cắt ngang"], "#EAF2F9", ft, fb)
    box(d, (665, 65, 1140, 265), "Robot và cảm biến", ["LiDAR, odometry, vị trí", "Trạng thái goal và va chạm"], "#EFF5EE", ft, fb)
    box(d, (1330, 65, 1765, 265), "SAC đã đóng băng", ["Quan sát mỗi khoảng 1 s", "Đề xuất speed scale"], "#F7F1E7", ft, fb)
    box(d, (1330, 435, 1765, 635), "Supervisor", ["Xuất Nav2 SpeedLimit", "Giới hạn cả vx và wz"], "#F7F1E7", ft, fb)
    box(d, (665, 435, 1140, 635), "Nav2 MPPI", ["Chọn quỹ đạo và cặp vx, wz", "Theo giới hạn đang áp dụng"], "#EAF2F9", ft, fb)
    box(d, (35, 435, 485, 635), "Chuyển động thực", ["Robot chạy hoặc dừng", "Ghi tốc độ, goal, va chạm"], "#EFF5EE", ft, fb)
    arrow(d, (485, 165), (665, 165))
    arrow(d, (1140, 165), (1330, 165))
    arrow(d, (902, 265), (902, 435))
    arrow(d, (1548, 265), (1548, 435))
    arrow(d, (1330, 535), (1140, 535))
    arrow(d, (665, 535), (485, 535))
    arrow(d, (260, 435), (260, 265))
    d.text((43, 700), "pure_mppi: bỏ nhánh SAC và supervisor; MPPI dùng giới hạn nền 0,70 m/s.", font=small, fill="#30465A")
    d.text((43, 743), "sac_only: bật nhánh SAC, tắt TTC override và front safety stop; checkpoint không học thêm.", font=small, fill="#30465A")
    DIAGRAM.parent.mkdir(parents=True, exist_ok=True)
    im.save(DIAGRAM)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), "D9D9D9")
        borders.append(tag)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    for i, header in enumerate(headers):
        table.columns[i].width = Inches(widths[i])
        cell = table.rows[0].cells[i]
        cell.text = header
        set_cell_shading(cell, "DCE9F4")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_border(cell)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row_idx, data in enumerate(rows):
        cells = table.add_row().cells
        for col, value in enumerate(data):
            cells[col].text = str(value)
            cells[col].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cells[col])
            if row_idx % 2:
                set_cell_shading(cells[col], "F4F8FB")
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(9)
    return table


def add_para(doc, text, bold_lead=None):
    p = doc.add_paragraph(style="Normal")
    if bold_lead:
        p.add_run(bold_lead).bold = True
    p.add_run(text)
    return p


def build():
    make_diagram()
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.68)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.77)
    sec.right_margin = Inches(0.77)
    styles = doc.styles
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = styles[name]
        style.font.name = "DejaVu Sans"
        style.font.color.rgb = RGBColor(0, 0, 0)
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Title"].font.size = Pt(18)
    styles["Title"].font.bold = True
    styles["Title"].paragraph_format.space_after = Pt(9)
    styles["Heading 1"].font.size = Pt(12)
    styles["Heading 1"].paragraph_format.space_before = Pt(12)
    styles["Heading 1"].paragraph_format.space_after = Pt(5)
    styles["Heading 2"].font.size = Pt(10.5)
    styles["Heading 2"].paragraph_format.space_before = Pt(9)
    styles["Heading 2"].paragraph_format.space_after = Pt(4)

    doc.add_paragraph("Sơ đồ đánh giá MPPI và SAC trên hai map warehouse", style="Title")
    add_para(doc, "Tài liệu này mô tả phép thử trên hai map warehouse và kết quả MPPI thuần ở trần 0,70 m/s. Trên cả hai map, vận tốc đoạn thoáng gần bằng trần nhưng tỷ lệ va chạm cao; đó là căn cứ để thử SAC giảm trần trước nguy cơ giao cắt, chưa phải bằng chứng rằng tăng trần chắc chắn kém hơn.")

    doc.add_paragraph("Luồng điều khiển đang được đánh giá", style="Heading 1")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(DIAGRAM), width=Inches(6.9))

    doc.add_paragraph("Một lượt thử diễn ra như thế nào", style="Heading 1")
    add_table(doc, ["Bước", "Cách chạy hiện tại"], [
        ("1  Khởi tạo", "Reset robot; giữ nguyên map, Nav2 và MPPI; gửi goal cố định tại (9, 0)."),
        ("2  Đoạn đầu", "Robot đi từ điểm xuất phát. Ghi tốc độ thực trên đoạn x từ 0,5 m tới trước ngưỡng kích hoạt vật cản."),
        ("3  Cắt ngang", "Khi robot đạt trigger x, vật cản động xuất hiện tại x khoảng 4,7–5,3 m rồi đi ngang với tốc độ 0,25–0,55 m/s."),
        ("4  Sau khi đi ngang", "Vật cản ở lại tại vị trí cuối cho đến hết lượt. Robot tiếp tục tới goal, va chạm hoặc timeout."),
        ("5  Kết thúc", "Ghi succeeded, collision hoặc timeout sau tối đa 75 s thời gian mô phỏng."),
    ], [1.18, 5.72])

    doc.add_paragraph("Các biến được thay đổi và giữ cố định", style="Heading 1")
    add_para(doc, "Mỗi map chạy 100 lượt với seed 42. Tình huống thay đổi theo vị trí cắt ngang, khoảng cách kích hoạt trước điểm cắt 1,7–2,8 m, phía xuất phát và tốc độ vật cản. Bản đồ, goal và trần MPPI nền 0,70 m/s được giữ cố định trong một phép so sánh. Danh sách tình huống có thể tái tạo bằng cùng seed hoặc phát lại chính xác theo trial ID.")
    add_para(doc, "Phạm vi tài liệu chỉ gồm mppi_warehouse_aisles và mppi_warehouse_mixed. Mỗi map có vật cản tĩnh trong bản đồ và một vật cản động cắt ngang trong từng lượt thử.")

    doc.add_page_break()
    doc.add_paragraph("Kết quả MPPI thuần trên hai map", style="Heading 1")
    add_para(doc, "Cùng giới hạn nền vx_max = 0,70 m/s, seed 42, 100 lượt mỗi map và timeout 75 s mỗi lượt. SAC và supervisor RL không hoạt động trong nhánh này.")
    add_table(doc, ["Map", "Tới goal", "Va chạm", "Timeout", "Va chạm khi vật còn chạy", "Tốc độ đoạn thoáng"], [
        ("Warehouse aisles", "7/100", "93/100", "0/100", "93/100 lượt gặp", "0,700 m/s"),
        ("Warehouse mixed", "13/100", "87/100", "0/100", "87/99 lượt gặp", "0,700 m/s"),
    ], [1.45, 0.72, 0.72, 0.72, 1.56, 1.73])
    add_para(doc, "Kết quả cho thấy điểm yếu quan sát được là va chạm trong tình huống giao cắt, không phải thiếu tốc độ ở đoạn thoáng. Vì vậy, hướng thử nghiệm tiếp theo là hạ trần tốc độ khi có nguy cơ và kiểm tra liệu tỷ lệ va chạm giảm mà không tăng timeout hoặc thời gian tới goal. Chưa có nhánh MPPI với trần cao hơn chạy trên đúng bộ tình huống này, nên không thể khẳng định bằng dữ liệu rằng tăng trần luôn kém hơn.")

    doc.add_paragraph("Phép đánh giá đo những gì", style="Heading 1")
    add_table(doc, ["Nhóm chỉ số", "Ý nghĩa và giới hạn"], [
        ("Kết quả nhiệm vụ", "Tỷ lệ tới goal, số va chạm, số timeout. Timeout có thể bao gồm kẹt, nhưng hiện không phải một nhãn deadlock riêng trong benchmark."),
        ("Hiệu suất", "Tốc độ thực trung bình ở đoạn trước khi vật cản xuất hiện; thời gian tới goal chỉ tính trên các lượt thành công."),
        ("Phản ứng", "Scale SAC yêu cầu và scale supervisor áp dụng; lệnh MPPI, lệnh cuối, vận tốc thực; độ trễ từ lúc front LiDAR báo gần tới lúc giảm lệnh hoặc giảm vận tốc."),
        ("Khoảng cách", "Khoảng cách tâm và độ hở thân xe ước tính theo ground truth; không thay thế cho sự kiện va chạm do mô phỏng báo."),
        ("Gặp vật đang chạy", "moving_encounter chỉ tính nếu robot tới cách tâm vật cản không quá 2 m khi vật cản còn di chuyển. Va chạm sau khi vật đứng lại vẫn tính vào tổng va chạm, nhưng không thuộc mẫu số này."),
    ], [1.45, 5.45])

    doc.add_paragraph("Cách đọc kết quả giữa các nhánh", style="Heading 1")
    add_table(doc, ["Nhánh", "Câu hỏi trả lời được"], [
        ("pure_mppi", "Nav2 MPPI nền hoàn thành hành trình thế nào khi không có supervisor RL?"),
        ("sac_only", "Hệ MPPI cộng checkpoint SAC cố định và supervisor SpeedLimit hoạt động thế nào khi không có TTC override và front safety stop?"),
        ("sac và shield_only", "Ảnh hưởng của các lớp supervisor hoặc safety; không dùng làm bằng chứng SAC tự xử lý vật cản."),
    ], [1.45, 5.45])

    doc.add_paragraph("Ranh giới kết luận hiện tại", style="Heading 1")
    add_para(doc, "Benchmark đo toàn hệ. SpeedLimit hiện tại giảm cả trần vận tốc thẳng vx và trần vận tốc quay wz; do đó một lượt kẹt quay không thể tự động quy cho SAC chọn tốc độ sai. Phép thử cũng chưa xác lập vận tốc tối ưu, chưa đánh giá mọi kiểu vật cản động, và chưa kiểm tra khả năng tổng quát với điểm đầu hoặc goal mới.")
    add_para(doc, "Khi so sánh tỷ lệ tránh vật cản đang cắt ngang, cần báo riêng số moving_encounter làm mẫu số. Một lượt thành công sau khi vật cản đã đứng lại chỉ chứng minh robot đi qua tình huống tổng thể, không chứng minh tránh được vật đang di chuyển.")

    doc.add_paragraph("Phép so sánh cần có trước quyết định tách giới hạn", style="Heading 1")
    add_para(doc, "Giữ cùng checkpoint SAC, map, seed và từng tình huống; chỉ thay cách áp giới hạn: SpeedLimit hiện tại giảm cả vx và wz so với cách chỉ giới hạn vx. Đối chiếu tỷ lệ tới goal, va chạm trong moving_encounter, timeout, thời gian tới goal và các đoạn robot quay tại chỗ. Đây là phép thử đề xuất, chưa phải kết quả đã có.")

    doc.add_paragraph("Căn cứ trong dự án", style="Heading 2")
    add_para(doc, "run_mppi_three_maps.py xác định các map, nhánh thử, 100 lượt mặc định, timeout và rosbag. benchmark_dynamic_obstacle.py tạo kịch bản, ghi mẫu thời gian và tổng hợp chỉ số. rl_velocity_supervisor.py triển khai scale qua Nav2 SpeedLimit. Số liệu MPPI thuần lấy từ benchmark_results/mppi_three_maps_20261011_235102 của hai map warehouse.")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
