#!/usr/bin/env python3
"""Create the Vietnamese speed-scale handoff note for teammates."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports" / "giai_thich_speed_scale_supervisor_cho_nhom.docx"


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def borders(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    elem = OxmlElement("w:tcBorders")
    for name in ("top", "left", "bottom", "right"):
        side = OxmlElement(f"w:{name}")
        side.set(qn("w:val"), "single")
        side.set(qn("w:sz"), "4")
        side.set(qn("w:color"), "D9D9D9")
        elem.append(side)
    tcpr.append(elem)


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for idx, head in enumerate(headers):
        t.columns[idx].width = Inches(widths[idx])
        c = t.rows[0].cells[idx]
        c.text = head
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(c, "DCE9F4")
        borders(c)
        for run in c.paragraphs[0].runs:
            run.bold = True
    for ri, row in enumerate(rows):
        cs = t.add_row().cells
        for ci, value in enumerate(row):
            cs[ci].text = str(value)
            cs[ci].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            borders(cs[ci])
            if ri % 2:
                shade(cs[ci], "F5F8FB")
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.12
                for run in p.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def paragraph(doc, text):
    return doc.add_paragraph(text, style="Normal")


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.68)
    sec.left_margin = Inches(0.78)
    sec.right_margin = Inches(0.78)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        st = doc.styles[name]
        st.font.name = "DejaVu Sans"
        st.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.14
    title = doc.styles["Title"]
    title.font.size = Pt(17)
    title.font.bold = True
    title.paragraph_format.space_after = Pt(10)
    for name, size in (("Heading 1", 12), ("Heading 2", 10.5)):
        st = doc.styles[name]
        st.font.size = Pt(size)
        st.paragraph_format.space_before = Pt(12)
        st.paragraph_format.space_after = Pt(5)

    doc.add_paragraph("Thiết kế SAC điều chỉnh trần tốc độ bằng speed scale", style="Title")
    paragraph(doc, "Ghi chú này giúp nhóm đọc đúng quyết định thiết kế hiện hành trên hai map warehouse. SAC chỉ đề xuất một hệ số giới hạn tốc độ; SAC không xuất cmd_vel và không sửa trực tiếp hai tham số vx_max, wz_max. Mục tiêu thử nghiệm là giảm tốc trước nguy cơ giao cắt mà vẫn duy trì tốc độ khi đường thoáng.")

    doc.add_paragraph("Vì sao chọn một biến speed scale", style="Heading 1")
    paragraph(doc, "Cho SAC xuất trực tiếp vx và wz sẽ biến bài toán thành học lái và học tránh vật cản đồng thời. Nhóm hiện muốn kiểm tra câu hỏi hẹp hơn: một chính sách học khi nào cần hạ trần tốc độ có giảm va chạm hay không. Action một chiều giúp tách quyết định điều tiết tốc độ khỏi quyết định hướng chuyển động. Ghi tham số vx_max, wz_max liên tục cũng không phải đường điều khiển hiện dùng; supervisor phát SpeedLimit để áp trần mà không gọi callback cập nhật tham số động của MPPI.")

    doc.add_paragraph("Speed scale đi qua supervisor như thế nào", style="Heading 1")
    p = doc.add_paragraph("Quan sát → SAC đề xuất scale → supervisor kiểm tra và làm mượt → SpeedLimit → trần hiệu dụng → lệnh vận tốc → vận tốc thực")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(9)
    paragraph(doc, "SAC nhận quan sát rồi đề xuất một số trong khoảng 0,15 đến 1,00 khoảng mỗi 1 giây. Supervisor nhận số đó qua /rl_velocity_limits, chặn giá trị ngoài khoảng, làm mượt theo thời gian và công bố giới hạn áp dụng trên /speed_limit. Nó cũng công bố xác nhận trên /rl_velocity_limits_applied. Vì có làm mượt, scale yêu cầu và scale đã áp dụng có thể khác nhau trong cùng một thời điểm.")

    doc.add_paragraph("Thông số hiện hành của nhánh SAC only", style="Heading 1")
    table(doc, ["Thông số", "Giá trị", "Ý nghĩa"], [
        ("Action SAC", "0,15 đến 1,00", "Chỉ hạ hoặc trả về trần nền; không tăng quá nền."),
        ("Chu kỳ action", "Khoảng 1 giây", "SAC đưa đề xuất mới mỗi bước đánh giá."),
        ("Supervisor", "5 Hz", "Cập nhật trần hiệu dụng giữa các action."),
        ("Làm mượt mặc định", "0,25 scale/giây; alpha 0,35", "Hạn chế thay đổi trần quá đột ngột."),
        ("Watchdog", "2,5 giây", "Mất action hợp lệ thì trả scale về 1,00."),
        ("Lớp an toàn riêng", "Tắt trong SAC only", "TTC override và front safety stop không can thiệp vào phép thử này."),
    ], [1.55, 1.42, 3.97])

    doc.add_paragraph("Scale hiện còn tác động đến vmax không", style="Heading 1")
    paragraph(doc, "Có. Nó không thay giá trị cấu hình nền vx_max = 0,70 m/s và wz_max = 0,50 rad/s, nhưng tạo trần hiệu dụng thấp hơn. Cơ chế đang chạy dùng một tỷ lệ cho cả vận tốc tiến và vận tốc quay.")
    table(doc, ["Scale áp dụng", "Trần tiến hiệu dụng", "Trần quay hiệu dụng"], [
        ("1,00", "0,70 m/s", "0,50 rad/s"),
        ("0,50", "0,35 m/s", "0,25 rad/s"),
        ("0,15", "0,105 m/s", "0,075 rad/s"),
    ], [1.42, 2.72, 2.80])
    paragraph(doc, "Scale 1,00 chỉ bỏ giới hạn bổ sung, không ra lệnh chạy ở 0,70 m/s. Scale 0,50 cũng không bảo đảm vận tốc thực bằng 0,35 m/s; lệnh thực có thể thấp hơn hoặc bằng 0. Scale tối thiểu 0,15 không phải lệnh dừng khẩn cấp. Việc chỉ giới hạn vận tốc tiến mà giữ nguyên khả năng quay vẫn là phương án đang xem xét, chưa được triển khai trong nhánh này.")

    doc.add_paragraph("Cách kiểm tra robot thật sự thay đổi tốc độ", style="Heading 1")
    paragraph(doc, "Benchmark ghi chuỗi mẫu theo thời gian gồm x, y, scale SAC yêu cầu, scale supervisor áp dụng, lệnh vận tốc trước và sau, vận tốc tiến và quay thực tế, cùng vị trí và trạng thái vật cản. Rosbag ghi lại các topic tương ứng. Vì vậy có thể đối chiếu bốn tầng: scale yêu cầu → scale áp dụng → lệnh vận tốc → vận tốc thực tại từng vị trí trên map. Báo cáo tổng hợp hiện có tốc độ trung bình đoạn thoáng, thời gian tới goal, va chạm và timeout; chưa có heatmap tốc độ theo từng vùng bản đồ.")
    paragraph(doc, "Ví dụ từ benchmark checkpoint SAC 20k: tốc độ thực trung bình ở đoạn thoáng là khoảng 0,333 m/s trên warehouse aisles và 0,512 m/s trên warehouse mixed. MPPI thuần ở cùng trần nền đạt khoảng 0,700 m/s trên cả hai map. Điều này xác nhận hệ thống đã làm robot chậm đi thật, nhưng cũng đặt câu hỏi liệu SAC có đang giảm tốc quá nhiều ngay cả khi đường thoáng. Không nên dùng riêng tỷ lệ va chạm để kết luận chính sách đã học đúng thời điểm giảm và tăng tốc.")

    doc.add_paragraph("Điểm cần giữ khi đọc kết quả", style="Heading 1")
    paragraph(doc, "Kết quả hiện tại là của toàn chuỗi SAC, supervisor và bộ điều khiển. Khi scale thấp, cả trần quay cũng thấp nên timeout hoặc kẹt quay không tự động chứng minh SAC chọn sai tốc độ tiến. Muốn kiểm tra thay đổi đề xuất, cần giữ cùng checkpoint và cùng tình huống, chỉ thay cách áp trần rồi so sánh vận tốc thực, va chạm, timeout và thời gian tới goal.")

    doc.add_paragraph("Tệp để teammate đối chiếu", style="Heading 2")
    paragraph(doc, "agv_rl_env.py: action một chiều và quan sát. rl_velocity_supervisor.py: kiểm tra, làm mượt, watchdog và phát SpeedLimit. rl_supervisor_sac_only.yaml: giá trị vận hành. benchmark_dynamic_obstacle.py: dữ liệu mẫu và chỉ số. run_mppi_three_maps.py: hai map được chọn và các topic rosbag. Kết quả ví dụ nằm trong benchmark_results/sac_only_conservative_two_warehouse_seed42_retry.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
